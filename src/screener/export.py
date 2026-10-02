# src/screener/export.py
"""Day 17: Excel export with per-cell threshold colour-coding."""

from pathlib import Path
import pandas as pd
from openpyxl.styles import PatternFill

GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

DISPLAY_COLUMNS = [
    "company_id", "broad_sector", "return_on_equity_pct", "debt_to_equity",
    "interest_coverage", "free_cash_flow_cr", "revenue_cagr_5yr", "pat_cagr_5yr",
    "operating_profit_margin_pct", "pe_ratio", "pb_ratio", "dividend_yield_pct",
    "sales", "net_profit", "eps_cagr_5yr", "asset_turnover", "market_cap_crore",
    "dividend_payout_ratio_pct", "net_profit_margin_pct", "cash_from_operations_cr",
    "composite_quality_score",
]


def _passes_filter(row: pd.Series, filt: dict) -> bool | None:
    """Returns True/False if evaluable, None if metric column missing (skip colouring)."""
    metric, operator_str, value = filt["metric"], filt["operator"], filt["value"]
    if metric not in row.index or pd.isna(row[metric]):
        return None
    val = row[metric]
    return {">": val > value, "<": val < value, ">=": val >= value,
            "<=": val <= value, "==": abs(val - value) < 1e-6}.get(operator_str)


def export_screener_output(results: dict, config: dict, output_path: str | Path) -> Path:
    """One sheet per preset, colour-coded cells (green=passes threshold, red=fails)."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for preset_name, df in results.items():
            cols = [c for c in DISPLAY_COLUMNS if c in df.columns]
            sheet_df = df[cols].copy()
            sheet_name = preset_name[:31]  # Excel sheet-name limit
            sheet_df.to_excel(writer, sheet_name=sheet_name, index=False)

            if preset_name == "Turnaround Watch":
                continue  # no static filter list to colour against

            filters = config["presets"].get(preset_name, [])
            ws = writer.sheets[sheet_name]
            for row_idx, (_, row) in enumerate(sheet_df.iterrows(), start=2):
                for filt in filters:
                    metric = filt["metric"]
                    if metric not in cols:
                        continue
                    col_idx = cols.index(metric) + 1
                    result = _passes_filter(row, filt)
                    if result is True:
                        ws.cell(row=row_idx, column=col_idx).fill = GREEN_FILL
                    elif result is False:
                        ws.cell(row=row_idx, column=col_idx).fill = RED_FILL

    return output_path