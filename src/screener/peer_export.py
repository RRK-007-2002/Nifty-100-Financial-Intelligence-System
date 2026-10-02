# src/analytics/peer_export.py
"""
Day 20: peer_comparison.xlsx — 11 sheets (one per peer group), percentile
colour-coding, benchmark-row highlight, median summary row.
Reuses Day 18's compute_peer_percentiles() output (long format) and
pivots it wide for display.
"""

from pathlib import Path
import pandas as pd
from openpyxl.styles import PatternFill

GREEN = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
YELLOW = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
RED = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
GOLD = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")


def _percentile_fill(pct: float):
    if pd.isna(pct):
        return None
    if pct >= 0.75:
        return GREEN
    if pct <= 0.25:
        return RED
    return YELLOW


# def _pivot_group(peer_percentiles: pd.DataFrame, group_name: str) -> pd.DataFrame:
#     """Long-format peer_percentiles -> wide table: one row/company, metric columns as percentiles."""
#     grp = peer_percentiles[
#         (peer_percentiles["peer_group_name"] == group_name) & (peer_percentiles["metric"] != "__no_group__")
#     ]
#     value_pivot = grp.pivot(index="company_id", columns="metric", values="value")
#     pct_pivot = grp.pivot(index="company_id", columns="metric", values="percentile_rank")
#     pct_pivot.columns = [f"{c}_percentile" for c in pct_pivot.columns]
#     return value_pivot.join(pct_pivot).reset_index()

def _pivot_group(peer_percentiles: pd.DataFrame, group_name: str) -> pd.DataFrame:
    """Long-format peer_percentiles -> wide table: one row/company, metric columns."""

    grp = peer_percentiles[
        (peer_percentiles["peer_group_name"] == group_name)
        & (peer_percentiles["metric"] != "__no_group__")
    ].copy()

    if grp.empty:
        return pd.DataFrame(columns=["company_id"])

    key_cols = ["company_id", "metric"]
    if "year" in grp.columns:
        grp = grp.sort_values(key_cols + ["year"], ascending=[True, True, False])
    grp = grp.drop_duplicates(subset=key_cols, keep="last")

    value_pivot = grp.pivot(index="company_id", columns="metric", values="value")
    pct_pivot = grp.pivot(index="company_id", columns="metric", values="percentile_rank")
    pct_pivot.columns = [f"{c}_percentile" for c in pct_pivot.columns]

    return value_pivot.join(pct_pivot).reset_index()
def export_peer_comparison(
    peer_percentiles: pd.DataFrame,
    companies: pd.DataFrame,
    peer_groups: pd.DataFrame,
    output_path: str | Path,
) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    group_names = sorted(peer_percentiles.loc[
        peer_percentiles["metric"] != "__no_group__", "peer_group_name"
    ].dropna().unique())

    if not group_names:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "No peer groups"
        ws.append(["company_id", "company_name", "note"])
        ws.append(["", "", "No peer group data available"])
        wb.save(output_path)
        return output_path

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for group_name in group_names:
            wide = _pivot_group(peer_percentiles, group_name)
            if wide.empty:
                continue
            wide = wide.merge(companies[["id", "company_name"]], left_on="company_id", right_on="id", how="left")
            wide = wide.drop(columns=["id"])
            cols = ["company_id", "company_name"] + [c for c in wide.columns if c not in ("company_id", "company_name")]
            wide = wide[cols]

            # Median summary row
            numeric_cols = wide.select_dtypes(include="number").columns
            median_row = {c: wide[c].median() for c in numeric_cols}
            median_row["company_id"] = "MEDIAN"
            median_row["company_name"] = ""
            wide = pd.concat([wide, pd.DataFrame([median_row])], ignore_index=True)

            sheet_name = group_name[:31]
            wide.to_excel(writer, sheet_name=sheet_name, index=False)
            ws = writer.sheets[sheet_name]

            if "is_benchmark" in peer_groups.columns:
                benchmark_ids = set(
                    peer_groups[(peer_groups["peer_group_name"] == group_name) & (peer_groups["is_benchmark"] == 1)]["company_id"]
                )
            else:
                benchmark_ids = set()

            pct_cols = [c for c in wide.columns if c.endswith("_percentile")]
            for row_idx, (_, row) in enumerate(wide.iterrows(), start=2):
                if row["company_id"] in benchmark_ids:
                    for col_idx in range(1, len(wide.columns) + 1):
                        ws.cell(row=row_idx, column=col_idx).fill = GOLD
                    continue  # benchmark highlight takes priority over percentile colour
                if row["company_id"] == "MEDIAN":
                    continue
                for col_name in pct_cols:
                    col_idx = wide.columns.get_loc(col_name) + 1
                    fill = _percentile_fill(row[col_name])
                    if fill:
                        ws.cell(row=row_idx, column=col_idx).fill = fill

    return output_path